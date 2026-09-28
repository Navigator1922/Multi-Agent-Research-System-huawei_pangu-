"""多格式文档加载、切分和 ChromaDB 向量入库。"""

from __future__ import annotations

import math
import re
import uuid
from pathlib import Path
from typing import Any, Iterable

from config.structure import ContextInput, DataBlock


_DEFAULT_EMBEDDING_MODEL = "BAAI/bge-small-zh-v1.5"
_CHUNK_MIN_LENGTH = 300
_CHUNK_MAX_LENGTH = 500

_IMAGE_SUFFIXES = {
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
_TEXT_SUFFIXES = {
    ".htm",
    ".html",
    ".json",
    ".jsonl",
    ".md",
    ".markdown",
    ".rst",
    ".txt",
    ".xml",
}
_TABLE_SUFFIXES = {".csv", ".xls", ".xlsx"}
_SUPPORTED_SUFFIXES = {
    ".docx",
    ".pdf",
    *_IMAGE_SUFFIXES,
    *_TEXT_SUFFIXES,
    *_TABLE_SUFFIXES,
}


def loader(data: Any) -> ContextInput:
    """解析文件或文件夹，并将文档块写入本地 ChromaDB。"""

    input_path = _resolve_input_path(data)
    files = _collect_files(input_path)
    blocks: list[DataBlock] = []

    for file_path in files:
        file_blocks = _parse_file(file_path, input_path)
        if not file_blocks:
            raise ValueError(f"文件没有解析出有效内容：{file_path}")
        blocks.extend(file_blocks)

    if not blocks:
        raise ValueError(f"输入路径没有可入库的内容：{input_path}")

    topic = _derive_topic(blocks)
    collection_id = _persist_to_chroma(blocks, input_path)
    source_id = _source_id_for_input(input_path)
    url = input_path.resolve().as_uri()

    return ContextInput(
        topic=topic,
        db_collection_id=collection_id,
        source_id=source_id,
        url=url,
    )


def _resolve_input_path(data: Any) -> Path:
    """将输入参数严格解析为一个存在的文件或文件夹路径。"""

    if isinstance(data, Path):
        input_path = data.expanduser()
    elif isinstance(data, str):
        input_path = Path(data).expanduser()
    elif isinstance(data, dict):
        path_value = data.get("path", data.get("file_path", data.get("folder_path")))
        if not isinstance(path_value, (str, Path)):
            raise TypeError("loader 的字典输入必须包含 path、file_path 或 folder_path")
        input_path = Path(path_value).expanduser()
    else:
        raise TypeError("loader 只接受文件路径、文件夹路径或包含路径字段的字典")

    if not input_path.exists():
        raise FileNotFoundError(f"输入路径不存在：{input_path}")
    if not input_path.is_file() and not input_path.is_dir():
        raise ValueError(f"输入路径既不是文件也不是文件夹：{input_path}")
    return input_path.resolve()


def _collect_files(input_path: Path) -> list[Path]:
    """收集并校验输入路径下的可解析文档。"""

    if input_path.is_file():
        files = [input_path]
    else:
        files = sorted(
            path
            for path in input_path.rglob("*")
            if path.is_file() and not path.name.startswith("~$")
        )

    if not files:
        raise ValueError(f"输入文件夹为空：{input_path}")

    unsupported = [
        path
        for path in files
        if path.suffix.lower() not in _SUPPORTED_SUFFIXES
    ]
    if unsupported:
        names = ", ".join(str(path) for path in unsupported[:10])
        suffix = "" if len(unsupported) <= 10 else " 等"
        raise ValueError(f"存在不支持的文件格式：{names}{suffix}")
    return files


def _parse_file(file_path: Path, input_root: Path) -> list[DataBlock]:
    """根据文件后缀调用对应解析器。"""

    suffix = file_path.suffix.lower()
    source_id = _source_id_for_file(file_path, input_root)
    base_metadata = {
        "source_id": source_id,
        "source_path": str(file_path),
        "file_name": file_path.name,
    }

    if suffix == ".pdf":
        return _parse_pdf(file_path, base_metadata)
    if suffix == ".docx":
        return _parse_docx(file_path, base_metadata)
    if suffix in _TABLE_SUFFIXES:
        return _parse_table_file(file_path, base_metadata)
    if suffix in _IMAGE_SUFFIXES:
        return [
            DataBlock(
                type="image",
                content=str(file_path),
                metadata={**base_metadata, "file_type": "image"},
            )
        ]
    if suffix in _TEXT_SUFFIXES:
        return _parse_text_file(file_path, base_metadata)
    raise ValueError(f"没有对应的文件解析器：{file_path}")


def _parse_pdf(file_path: Path, base_metadata: dict[str, Any]) -> list[DataBlock]:
    """使用 PyPDF2 逐页提取 PDF 文本并切分。"""

    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise RuntimeError("解析 PDF 需要安装 PyPDF2") from exc

    try:
        reader = PdfReader(str(file_path))
    except Exception as exc:
        raise RuntimeError(f"PDF 打开失败：{file_path}") from exc

    blocks: list[DataBlock] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise RuntimeError(f"PDF 第 {page_number} 页解析失败：{file_path}") from exc
        blocks.extend(
            _text_blocks(
                text,
                {**base_metadata, "page": page_number, "file_type": "pdf"},
            )
        )

    if not blocks:
        raise ValueError(f"PDF 未提取到文本内容，可能是扫描图片：{file_path}")
    return blocks


def _parse_docx(file_path: Path, base_metadata: dict[str, Any]) -> list[DataBlock]:
    """使用 python-docx 提取段落和表格。"""

    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("解析 DOCX 需要安装 python-docx") from exc

    try:
        document = Document(str(file_path))
    except Exception as exc:
        raise RuntimeError(f"DOCX 打开失败：{file_path}") from exc

    blocks: list[DataBlock] = []
    paragraph_text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )
    blocks.extend(
        _text_blocks(
            paragraph_text,
            {**base_metadata, "file_type": "docx"},
        )
    )

    for table_number, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        markdown = _rows_to_markdown(rows)
        for chunk_number, chunk in enumerate(_chunk_text(markdown)):
            blocks.append(
                DataBlock(
                    type="table",
                    content=chunk,
                    metadata={
                        **base_metadata,
                        "file_type": "docx",
                        "table": table_number,
                        "chunk": chunk_number,
                    },
                )
            )
    return blocks


def _parse_table_file(file_path: Path, base_metadata: dict[str, Any]) -> list[DataBlock]:
    """使用 pandas 读取 CSV 或 Excel，并转换为 Markdown 表格。"""

    try:
        import pandas as pd
    except ImportError as exc:
        raise RuntimeError("解析 CSV 或 Excel 需要安装 pandas") from exc

    try:
        if file_path.suffix.lower() == ".csv":
            tables = {"csv": pd.read_csv(file_path)}
        else:
            tables = pd.read_excel(file_path, sheet_name=None)
    except Exception as exc:
        raise RuntimeError(f"表格文件读取失败：{file_path}") from exc

    blocks: list[DataBlock] = []
    for sheet_name, table in tables.items():
        if table is None:
            continue
        rows = [list(table.columns)] + table.fillna("").astype(str).values.tolist()
        markdown = _rows_to_markdown(rows)
        for chunk_number, chunk in enumerate(_chunk_text(markdown)):
            blocks.append(
                DataBlock(
                    type="table",
                    content=chunk,
                    metadata={
                        **base_metadata,
                        "file_type": file_path.suffix.lower().lstrip("."),
                        "sheet": str(sheet_name),
                        "chunk": chunk_number,
                    },
                )
            )
    if not blocks:
        raise ValueError(f"表格文件没有有效数据：{file_path}")
    return blocks


def _parse_text_file(file_path: Path, base_metadata: dict[str, Any]) -> list[DataBlock]:
    """读取普通文本文件并切分为文本块。"""

    raw = file_path.read_bytes()
    text = ""
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        raise ValueError(f"文本文件编码无法识别：{file_path}")
    blocks = _text_blocks(
        text,
        {**base_metadata, "file_type": file_path.suffix.lower().lstrip(".")},
    )
    if not blocks:
        raise ValueError(f"文本文件为空：{file_path}")
    return blocks


def _text_blocks(text: str, metadata: dict[str, Any]) -> list[DataBlock]:
    """清理文本并按目标长度生成 DataBlock。"""

    normalized = re.sub(r"[ \t\r\f\v]+", " ", text)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    return [
        DataBlock(
            type="text",
            content=chunk,
            metadata={**metadata, "chunk": chunk_number},
        )
        for chunk_number, chunk in enumerate(_chunk_text(normalized))
    ]


def _chunk_text(text: str) -> list[str]:
    """将文本切分为长度尽量位于 300 至 500 字的连续块。"""

    text = text.strip()
    if not text:
        return []
    if len(text) <= _CHUNK_MAX_LENGTH:
        return [text]

    chunk_count = max(2, math.ceil(len(text) / 450))
    base_length, remainder = divmod(len(text), chunk_count)
    chunks = []
    start = 0
    for index in range(chunk_count):
        length = base_length + (1 if index < remainder else 0)
        end = start + length
        chunks.append(text[start:end])
        start = end
    return chunks


def _rows_to_markdown(rows: Iterable[Iterable[Any]]) -> str:
    """将二维表格数据转换为 Markdown 表格。"""

    normalized_rows = [
        [str(cell).replace("|", "\\|").replace("\n", " ") for cell in row]
        for row in rows
    ]
    normalized_rows = [row for row in normalized_rows if row]
    if not normalized_rows:
        return ""
    width = max(len(row) for row in normalized_rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in normalized_rows]
    lines = [
        "| " + " | ".join(normalized_rows[0]) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized_rows[1:])
    return "\n".join(lines)


def _derive_topic(blocks: list[DataBlock]) -> str:
    """从第一个文本或表格块的前 100 个字中提取主题。"""

    for block in blocks:
        if block.type not in {"text", "table"}:
            continue
        sample = re.sub(r"\s+", " ", block.content).strip()[:100]
        if sample:
            sentence = re.split(r"[。！？!?；;]", sample, maxsplit=1)[0].strip()
            return (sentence or sample)[:100]

    source_name = str(blocks[0].metadata.get("file_name", "图片资料"))
    return f"{source_name}相关资料"


def _persist_to_chroma(blocks: list[DataBlock], input_path: Path) -> str:
    """创建随机 Chroma 集合，并将全部 DataBlock 持久化。"""

    try:
        import chromadb
        from chromadb.utils import embedding_functions
    except ImportError as exc:
        raise RuntimeError(
            "向量入库需要安装 chromadb 及其依赖；请先执行 pip install chromadb"
        ) from exc

    model_name = _embedding_model_name()
    try:
        embedding_function = embedding_functions.SentenceTransformerEmbeddingFunction(
            model_name=model_name
        )
    except Exception as exc:
        raise RuntimeError(f"中文嵌入模型初始化失败：{model_name}") from exc

    chroma_path = Path("./chroma_db").resolve()
    try:
        client = chromadb.PersistentClient(path=str(chroma_path))
        collection_id = uuid.uuid4().hex
        collection = client.create_collection(
            name=collection_id,
            embedding_function=embedding_function,
            metadata={
                "source_path": str(input_path),
                "embedding_model": model_name,
            },
        )
    except Exception as exc:
        raise RuntimeError(f"ChromaDB 集合创建失败：{chroma_path}") from exc

    ids = []
    documents = []
    metadatas = []
    for block in blocks:
        block_id = uuid.uuid4().hex
        document = block.content
        if block.type == "image":
            document = f"图片文件路径：{block.content}"
        metadata = _chroma_metadata(block.metadata, block.type)
        ids.append(block_id)
        documents.append(document)
        metadatas.append(metadata)

    try:
        collection.add(ids=ids, documents=documents, metadatas=metadatas)
    except Exception as exc:
        raise RuntimeError(f"文档块写入 ChromaDB 失败：{collection_id}") from exc
    return collection_id


def _embedding_model_name() -> str:
    """读取嵌入模型配置，没有配置时使用中文默认模型。"""

    import os

    return os.getenv("CHROMA_EMBEDDING_MODEL", _DEFAULT_EMBEDDING_MODEL).strip() or _DEFAULT_EMBEDDING_MODEL


def _chroma_metadata(metadata: dict[str, Any], block_type: str) -> dict[str, Any]:
    """筛选 Chroma 支持的元数据类型，并标记块类型。"""

    result: dict[str, Any] = {"type": block_type}
    for key, value in metadata.items():
        if value is None:
            continue
        if isinstance(value, (str, int, float, bool)):
            result[str(key)] = value
        else:
            result[str(key)] = str(value)
    return result


def _source_id_for_file(file_path: Path, input_root: Path) -> str:
    """根据文件相对路径生成稳定且适合元数据使用的来源编号。"""

    if input_root.is_file():
        relative_name = file_path.name
    else:
        relative_name = file_path.relative_to(input_root).as_posix()
    source_id = re.sub(r"[^0-9A-Za-z_\u4e00-\u9fff-]+", "_", relative_name)
    source_id = re.sub(r"_+", "_", source_id).strip("_.-")
    return source_id or uuid.uuid4().hex


def _source_id_for_input(input_path: Path) -> str:
    """生成返回给 ContextInput 的输入来源编号。"""

    if input_path.is_file():
        return _source_id_for_file(input_path, input_path)
    source_id = re.sub(r"[^0-9A-Za-z_\u4e00-\u9fff-]+", "_", input_path.name)
    return re.sub(r"_+", "_", source_id).strip("_.-") or uuid.uuid4().hex


__all__ = ["loader"]
